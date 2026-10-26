# -*- coding: utf-8 -*-
"""Bangun draft Process Specification M06 (docx + xlsx + gambar decision tree) dari m06_data.py.

Pakai:  python manajemen/generator/build_m06.py
Hasil:  progress/PRPL_M06_2B_RajinBelajar_ProcessSpecification_draft.docx
        progress/PRPL_M06_2B_RajinBelajar_DecisionTable.xlsx
        progress/PRPL_M06_2B_RajinBelajar_DecisionTree_3.4.png
"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from m06_data import *

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "progress")
F_DOCX = os.path.join(OUT, "PRPL_M06_2B_RajinBelajar_ProcessSpecification_draft.docx")
F_XLSX = os.path.join(OUT, "PRPL_M06_2B_RajinBelajar_DecisionTable.xlsx")
F_TREE = os.path.join(OUT, "PRPL_M06_2B_RajinBelajar_DecisionTree_3.4.png")


# ---------------------------------------------------------------- decision tree (PyMuPDF)
def gambar_tree():
    import pymupdf
    n = len(DTREE["langkah"])
    bw, bh, gap = 250, 44, 34
    W, H = 640, 70 + (n + 1) * (bh + gap)
    doc = pymupdf.open(); pg = doc.new_page(width=W, height=H)
    biru, merah, hijau, hitam = (0.85, 0.9, 1), (1, 0.88, 0.88), (0.86, 0.96, 0.86), (0, 0, 0)

    def kotak(x, y, w, h, teks, warna):
        r = pymupdf.Rect(x, y, x + w, y + h)
        pg.draw_rect(r, color=hitam, fill=warna, width=0.8)
        pg.insert_textbox(r + (4, 5, -4, -2), teks, fontsize=8.5, align=1)

    pg.insert_text((20, 28), DTREE["judul"], fontsize=11)
    x0, y = 30, 50
    for i, (tanya, cab, hasil) in enumerate(DTREE["langkah"]):
        kotak(x0, y, bw, bh, tanya, biru)
        kotak(x0 + bw + 90, y, 250, bh, hasil, merah)
        pg.draw_line((x0 + bw, y + bh / 2), (x0 + bw + 90, y + bh / 2), color=hitam, width=0.8)
        pg.insert_text((x0 + bw + 30, y + bh / 2 - 4), cab, fontsize=8)
        pg.draw_line((x0 + bw / 2, y + bh), (x0 + bw / 2, y + bh + gap), color=hitam, width=0.8)
        pg.insert_text((x0 + bw / 2 + 5, y + bh + gap / 2 + 3), "Ya", fontsize=8)
        y += bh + gap
    kotak(x0, y, bw + 90 + 250, bh, DTREE["berhasil"], hijau)
    pg.get_pixmap(dpi=160).save(F_TREE)


# ---------------------------------------------------------------- docx
def tulis_docx():
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.section import WD_ORIENT, WD_SECTION

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)

    def tabel(header, rows, size=9, widths=None):
        t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = h
            for r in c.paragraphs[0].runs:
                r.bold = True; r.font.size = Pt(size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        doc.add_paragraph()
        return t

    def kode(teks):
        for baris in teks.split("\n"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            r = p.add_run(baris if baris else " "); r.font.name = "Consolas"; r.font.size = Pt(8.5)
        doc.add_paragraph()

    def lanskap(on):
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        if on and s.orientation != WD_ORIENT.LANDSCAPE or (not on and s.orientation == WD_ORIENT.LANDSCAPE):
            s.orientation = WD_ORIENT.LANDSCAPE if on else WD_ORIENT.PORTRAIT
            s.page_width, s.page_height = s.page_height, s.page_width

    doc.add_heading("Menginventarisasi Proses Primitif", 1)
    tabel(["No. Proses", "Nama Proses", "Input", "Output", "Store Dibaca atau Ditulis", "Event Terkait", "Bentuk Logika"],
          INVENTARIS, 8.5, [1.3, 2.6, 2.6, 2.6, 3.6, 1.6, 2.4])
    doc.add_paragraph(INDUK)

    doc.add_heading("Parameter dan Aturan Bisnis", 1)
    tabel(["Parameter", "Nilai", "Sumber"], PARAMETER, 9, [4, 5, 7.5])
    tabel(["ID", "Aturan Bisnis", "Sumber"], ATURAN, 8.5, [1.4, 9.5, 5.6])

    doc.add_heading("Memilih Bentuk Logika", 1)
    tabel(["No. Proses", "Karakter Logika", "Bentuk Logika", "Alasan Pemilihan"], TEKNIK, 8.5, [1.3, 6, 3, 6.2])

    doc.add_heading("Menulis Process Specification", 1)
    dt = {d["id"]: d for d in DT}
    for s in SPEC:
        doc.add_heading(f"Proses {s['no']} {s['nama']}", 2)
        tabel(["Elemen", "Isi"], [
            ("Nomor dan nama proses", f"{s['no']} {s['nama']}"), ("Tujuan", s["tujuan"]), ("Input", s["input"]),
            ("Output", s["output"]), ("Data store dibaca", s["baca"]), ("Data store ditulis", s["tulis"]),
            ("Precondition", s["pre"]), ("Postcondition", s["post"]), ("Aturan bisnis", s["aturan"]),
            ("Exception", s["exc"]), ("Referensi data", s["ref"]), ("Bentuk logika dan alasan", s["bentuk"]),
        ], 9, [4, 12.5])
        doc.add_paragraph("Structured English:").runs[0].bold = True
        kode(s["se"])
        for d in DT:
            if (s["no"] == "1" and d["id"] == "DT-1") or (s["no"] == "3.1" and d["id"] == "DT-2"):
                doc.add_paragraph(d["judul"]).runs[0].bold = True
                rows = [(k,) + tuple(v) for k, v in d["kondisi"]] + [(a,) + tuple(v) for a, v in d["tindakan"]]
                rows.append(("Keterangan gagal",) + tuple(d["ket"]))
                tabel(["Kondisi atau Tindakan"] + d["kolom"], rows, 8.5)
                doc.add_paragraph("Pemeriksaan kelengkapan: " + d["cek"])
        if s["no"] == "3.4":
            doc.add_paragraph(DTREE["judul"]).runs[0].bold = True
            doc.add_picture(F_TREE, width=Cm(16))

    lanskap(True)
    doc.add_heading("Memvalidasi dengan Contoh Uji", 1)
    doc.add_paragraph("Tanggal hari ini pada semua kasus uji = 2026-10-26 (Senin). Hasil penelusuran diperoleh dengan menjalankan "
                      "Structured English / decision table / decision tree secara manual terhadap data uji.")
    tabel(["ID Uji", "Proses", "Aturan atau Jalur", "Data Uji", "Hasil yang Diharapkan", "Hasil Penelusuran", "Status", "Temuan"],
          UJI, 7.5, [1.3, 1.1, 2.6, 5.2, 5.2, 4, 1.3, 3.3])

    doc.add_heading("Traceability Review", 1)
    tabel(["Event ID", "Process ID", "Input atau Output", "Entri Kamus Data", "Aturan Bisnis", "Kasus Uji", "Status"],
          TRACE, 8, [2, 1.8, 5, 4, 3, 3, 5])

    lanskap(False)
    doc.add_heading("Revision History", 1)
    tabel(["Versi", "Tanggal", "Artefak atau Bagian", "Perubahan", "Alasan", "Penyusun", "Reviewer"], REVISI, 8.5)
    doc.add_heading("Log Kontribusi Individu", 1)
    tabel(["Nama atau NIM", "Peran", "Aktivitas", "Durasi", "Artefak atau Bukti", "Refleksi Singkat"],
          [(a, "", "", "", "", "") for a in ANGGOTA], 8.5)
    doc.add_heading("Pertanyaan Refleksi", 1)
    for i, (q, a) in enumerate(REFLEKSI, 1):
        doc.add_paragraph(f"{i}. {q}").runs[0].bold = True
        doc.add_paragraph(a)
    doc.save(F_DOCX)


# ---------------------------------------------------------------- xlsx
def tulis_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = Workbook(); thin = Side(style="thin", color="999999"); head = PatternFill("solid", fgColor="DDEBF7")

    def sheet(ws, judul, header, rows, widths):
        ws.append([judul]); ws["A1"].font = Font(bold=True, size=12); ws.append([])
        ws.append(header)
        for c in ws[3]:
            c.font = Font(bold=True); c.fill = head
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows(min_row=3):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

    first = True
    for d in DT:
        ws = wb.active if first else wb.create_sheet()
        ws.title = d["id"]; first = False
        rows = [(k,) + tuple(v) for k, v in d["kondisi"]] + [(a,) + tuple(v) for a, v in d["tindakan"]]
        rows += [("Keterangan gagal",) + tuple(d["ket"]), (), ("Pemeriksaan kelengkapan", d["cek"])]
        sheet(ws, d["judul"], ["Kondisi atau Tindakan"] + d["kolom"], rows, [48] + [16] * len(d["kolom"]))
    sheet(wb.create_sheet("DTree-1"), DTREE["judul"], ["Urutan", "Pertanyaan", "Bila Tidak", "Bila Ya"],
          [(i, q, h, "lanjut ke pertanyaan berikutnya" if i < len(DTREE["langkah"]) else DTREE["berhasil"])
           for i, (q, _, h) in enumerate(DTREE["langkah"], 1)], [8, 45, 45, 55])
    sheet(wb.create_sheet("Paket_Uji"), "Paket Uji", ["ID Uji", "Proses", "Aturan atau Jalur", "Data Uji", "Hasil yang Diharapkan",
                                                     "Hasil Penelusuran", "Status", "Temuan"], UJI, [8, 7, 22, 45, 45, 35, 8, 30])
    sheet(wb.create_sheet("Traceability"), "Traceability Review", ["Event ID", "Process ID", "Input atau Output", "Entri Kamus Data",
                                                                   "Aturan Bisnis", "Kasus Uji", "Status"], TRACE, [12, 10, 38, 28, 20, 18, 34])
    wb.save(F_XLSX)


if __name__ == "__main__":
    gambar_tree(); tulis_docx(); tulis_xlsx()
    print("Ditulis:\n ", F_DOCX, "\n ", F_XLSX, "\n ", F_TREE)
