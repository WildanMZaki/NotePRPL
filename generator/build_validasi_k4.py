# -*- coding: utf-8 -*-
"""Dokumen validasi silang Kelompok 6 terhadap M04 (DFD) & M05 (Kamus Data) Kelompok 4.

Pakai: python manajemen/generator/build_validasi_k4.py
Acuan: Modul 4 §1.4-1.6, §3.5, Lampiran C; Modul 5 §1.3, §1.6-1.11, Langkah 5.
"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "validasi-silang", "PRPL_2B_RajinBelajar_ValidasiSilang_Kelompok4_M04-M05.docx")

IDENTITAS = [
    ("Kelompok peninjau", "Kelompok 6 (RajinBelajar) - 2B"),
    ("Kelompok ditinjau", "Kelompok 4 - 2B"),
    ("Artefak", "M04 DFD Level 0 & Level 1 (gambar 'Sumber DFD' pada dokumen M05 Kelompok 4); "
                "M05 PRPL_M05_2B_Kelompok 4_DataDictionary.docx"),
    ("Acuan pemeriksaan", "Modul 4: §1.4 Balancing, §1.5 Audit sintaks-semantik-visual, §1.6 Black hole/Miracle/Gray hole, "
                          "§1.8 DFD logis, §3.5 Peer review, Lampiran C. Modul 5: §1.3, §1.6-§1.11, Langkah Kerja 5 Validasi Silang."),
    ("Tanggal", "2026-10-21"),
    ("Versi", "1.0"),
]

# --------------------------------------------------------------- Modul 4
M04_REVIEW = [
    # diagram, interpretasi, sintaks, semantik, balancing, keterbacaan, catatan
    ("DFD Level 0",
     "Sistem menerima pendaftaran anggota (validasi ke Sistem Akademik), memproses peminjaman dari anggota, "
     "memproses pengembalian & denda oleh petugas, dan menyusun laporan bulanan untuk Kepala Perpustakaan.",
     "Tidak", "Ya", "-", "Ya",
     "Sintaks: semua aliran ke/dari data store tidak bernama (VS-01). Proses 4 tidak menerima input dari Kepala "
     "Perpustakaan; sah bila laporan dipicu waktu, perlu klarifikasi (VS-07)."),
    ("DFD Level 1 Proses 2 Kelola Peminjaman",
     "2.1 memvalidasi kelayakan anggota, 2.2 memeriksa ketersediaan buku fisik, 2.3 mencatat transaksi dan "
     "menerbitkan bukti peminjaman untuk anggota.",
     "Tidak", "Tidak", "Ya", "Ya",
     "Semantik: tidak ada output ke Anggota bila anggota tidak valid (2.1) atau buku tidak tersedia (2.2); jalur gagal "
     "berakhir tanpa output (VS-03). Balancing: dibanding review sebelumnya, L0 sekarang sudah dua arah ke Data Koleksi "
     "sehingga sudah balance."),
    ("DFD Level 1 Proses 3 Kelola Pengembalian",
     "3.1 memeriksa tanggal pengembalian, 3.2 menghitung denda bila terlambat, 3.3 mencatat pengembalian dan "
     "menerbitkan bukti, 3.4 memperbarui status koleksi menjadi tersedia.",
     "Tidak", "Tidak", "Tidak", "Ya",
     "Semantik: gray hole di 3.2 dan 3.3 (VS-04, VS-05). Balancing: L0 Proses 3 dua arah ke Data Koleksi, tetapi di "
     "anak hanya ada tulis dari 3.4 (VS-02). Penamaan triggerStatusBuku (VS-06)."),
]

BALANCING = [
    # induk, arah, anak, status, catatan
    ("P2: pengajuanPeminjaman (Anggota -> 2)", "Masuk", "pengajuanPeminjaman (Anggota -> 2.1)", "Balance", "-"),
    ("P2: buktiPeminjaman (2 -> Anggota)", "Keluar", "buktiPeminjaman (2.3 -> Anggota)", "Balance", "-"),
    ("P2: D1 -> 2 (tanpa nama)", "Masuk", "D1 -> 2.1 (tanpa nama)", "Perlu klarifikasi", "Arah sama, nama tidak ada sehingga makna tidak bisa dibandingkan"),
    ("P2: D2 <-> 2 (tanpa nama)", "Dua arah", "D2 <-> 2.2 (tanpa nama)", "Perlu klarifikasi", "Arah sama; nama tidak ada"),
    ("P2: 2 -> D3 (tanpa nama)", "Keluar", "2.3 -> D3 (tanpa nama)", "Perlu klarifikasi", "Arah sama; nama tidak ada"),
    ("P3: dataPengembalianBuku (Petugas -> 3)", "Masuk", "dataPengembalianBuku (Petugas -> 3.1)", "Balance", "-"),
    ("P3: buktiPengembalian (3 -> Petugas)", "Keluar", "buktiPengembalian (3.3 -> Petugas)", "Balance", "-"),
    ("P3: infoDenda (3 -> Petugas)", "Keluar", "infoDenda (3.2 -> Petugas)", "Balance", "-"),
    ("P3: D2 <-> 3 (tanpa nama)", "Dua arah", "3.4 -> D2 (tulis saja)", "Tidak balance", "Sisi baca D2 pada induk tidak muncul di anak (VS-02)"),
    ("P3: D3 <-> 3 (tanpa nama)", "Dua arah", "D3 -> 3.1 (baca), 3.3 -> D3 (tulis)", "Perlu klarifikasi", "Arah sesuai; nama tidak ada"),
]

M05_CHECK = [
    ("Seluruh data flow dan data store pada DFD tercatat dalam inventaris", "Tidak",
     "Aliran ke/dari data store tidak tercatat karena tidak bernama di DFD (VS-10)."),
    ("Tidak ada sinonim atau homonim tanpa pengendalian nama", "Tidak",
     "StatusAnggota / statusAnggota / StatusKeanggotaan; IDTransaksi / IDPeminjaman; daftar alias tidak ada (VS-15)."),
    ("Setiap struktur komposit berakhir pada elemen elementer", "Ya",
     "Hampir semua struktur diurai sampai angka/huruf. Kekurangan kecil di VS-21."),
    ("Definisi menjelaskan makna bisnis, bukan mengulang nama", "Ya", "Definisi elemen sudah bermakna bisnis."),
    ("Format, domain, dan contoh data tidak saling bertentangan", "Tidak", "Lihat VS-14, VS-16, VS-17."),
    ("Setiap kunci dapat diidentifikasi dan aturan null dinyatakan", "Sebagian",
     "@ dipakai ganda dan di flow; struktur Data Transaksi tanpa @; beberapa elemen tanpa nullability (VS-12, VS-20)."),
    ("TBD memiliki pemilik serta rencana penyelesaian", "Tidak", "Volume = TBD tanpa pemilik/batas waktu; tidak ada daftar TBD (VS-18)."),
]

TRACE = [
    # langkah, hasil
    ("1. Telusuri input-output Proses 3 Kelola Pengembalian",
     "Input dataPengembalianBuku: ada entri. Output infoDenda dan buktiPengembalian: ada entri. Internal statusKeterlambatan dan "
     "triggerStatusBuku: ada entri. Aliran D3 -> 3.1, 3.3 -> D3, dan 3.4 -> D2: TIDAK ADA entri karena tidak bernama."),
    ("2. Elemen pada calon process spec terdefinisi",
     "Untuk menghitung TerlambatHari, 3.2 membutuhkan TanggalJatuhTempo dan TanggalKembali, tetapi statusKeterlambatan hanya "
     "[Terlambat | Tepat waktu]. TarifDendaPerHari dipakai di aturan, tetapi tidak didefinisikan sebagai elemen."),
    ("3. Sinonim, singkatan, homonim, kapitalisasi",
     "TotalDenda (infoDenda) vs TotalDendaBulanan (entri laporan) vs TotalDenda 1{angka}6 / 1{angka}8. "
     "StatusTransaksi [Dipinjam | Selesai] vs StatusPeminjaman [Tersedia | Tidak Tersedia]."),
    ("4. Store punya struktur record & kunci",
     "Data Anggota @IDAnggota dan Data Koleksi @KodeBuku jelas. Data Transaksi: entri komposit tanpa '=' dan penguraian tanpa '@'."),
    ("5. Masalah yang berasal dari DFD",
     "Beri nama semua aliran data store. Teruskan data tanggal ke 3.2 dan data buku ke 3.3, atau beri akses baca D3 (VS-04, VS-05)."),
    ("6. TBD", "Volume semua entri = TBD tanpa pemilik dan tenggat."),
]

TEMUAN = [
    # id, artefak, lokasi, jenis, deskripsi, severitas, rekomendasi
    ("VS-01", "M04", "L0, L1-P2, L1-P3: semua aliran data store", "Sintaks",
     "Seluruh aliran antara proses dan data store tidak memiliki nama. Acuan: Modul 1-3 tabel komponen DFD (Data Flow: 'kata benda "
     "atau struktur data', berlaku untuk semua aliran) dan Modul 5 §1.11 (seluruh data flow tercatat di inventaris; aliran tanpa nama tidak "
     "dapat diinventarisasi).", "Tinggi",
     "Beri nama kata benda spesifik pada setiap aliran (mis. dataAnggotaAktif, statusKoleksi, dataTransaksiBaru), lalu masukkan ke inventaris M05."),
    ("VS-02", "M04", "L0 Proses 3 <-> D2; L1-P3 3.4 -> D2", "Balancing",
     "Di induk, aliran Proses 3 dan Data Koleksi dua arah; di diagram anak hanya ada tulis dari 3.4. Sisi baca hilang.", "Sedang",
     "Jadikan aliran L0 searah (tulis) atau tambahkan aliran baca D2 pada proses anak yang membutuhkan."),
    ("VS-03", "M04", "L1-P2: 2.1, 2.2", "Semantik (jalur gagal)",
     "Bila anggota tidak valid (2.1) atau buku tidak tersedia (2.2), tidak ada output ke Anggota. Entri konfirmasiPeminjaman menyebut "
     "'mengembalikan nilai kosong'. Jalur gagal berakhir tanpa output.", "Tinggi",
     "Tambahkan output status/penolakan ke Anggota (dan padanannya di L0), lalu definisikan strukturnya di M05."),
    ("VS-04", "M04 + M05", "L1-P3: 3.2 Hitung Denda", "Semantik (gray hole)",
     "3.2 menghasilkan TerlambatHari dan TotalDenda, tetapi inputnya hanya statusKeterlambatan = [Terlambat | Tepat waktu] dan 3.2 tidak "
     "membaca store. Output tidak dapat dijelaskan dari input.", "Tinggi",
     "Perluas statusKeterlambatan agar membawa IDTransaksi + TanggalJatuhTempo + TanggalKembali, atau beri 3.2 akses baca D3."),
    ("VS-05", "M04 + M05", "L1-P3: 3.3 Catat Pengembalian", "Semantik (black/gray hole)",
     "Input 3.3 hanya infoDenda, yang menurut M05 'hanya muncul jika terlambat'. Pada pengembalian tepat waktu 3.3 tidak menerima input "
     "apa pun. buktiPengembalian juga memuat KodeBuku, JudulBuku, TanggalKembali yang tidak tersedia di input 3.3.", "Tinggi",
     "Alirkan data pengembalian dari 3.1 ke 3.3 untuk semua kasus (bukan hanya terlambat), atau beri 3.3 akses baca store yang diperlukan."),
    ("VS-06", "M04", "L1-P3: 3.3 -> 3.4 triggerStatusBuku", "Penamaan (DFD logis)",
     "'trigger' menggambarkan sinyal kontrol, bukan informasi. Modul 4 §1.8: nama aliran menjelaskan informasi.", "Sedang",
     "Ganti dengan kata benda informasi, mis. statusKoleksiKembali."),
    ("VS-07", "M04", "L0 Proses 4 Susun Laporan", "Perlu klarifikasi",
     "Proses 4 tidak menerima input dari Kepala Perpustakaan. Sah bila laporan dipicu waktu (temporal), tetapi tidak dijelaskan.", "Rendah",
     "Nyatakan event temporal di event table; bila laporan juga bisa diminta, tambahkan aliran permintaan laporan."),
    ("VS-08", "M04", "L0", "Keterbacaan",
     "Garis dari D1, D2, D3 menuju Proses 4 panjang dan bersilangan dengan aliran Proses 2/3.", "Rendah",
     "Pindahkan Proses 4 atau gunakan simbol duplikat store di dekat Proses 4."),
    ("VS-09", "M05", "Inventaris No 9 & 14 (infoDenda)", "Konsistensi istilah",
     "Satu nama infoDenda dipakai untuk dua aliran (3.2 -> Petugas dan 3.2 -> 3.3) dan berstatus 'Sinonim'. Ini nama yang sama, bukan "
     "sinonim; strukturnya harus identik.", "Rendah",
     "Pastikan kedua aliran memiliki isi yang sama. Ganti status entri menjadi 'Baku' atau bedakan namanya bila isinya berbeda."),
    ("VS-10", "M05", "Inventaris Data", "Kelengkapan",
     "Aliran ke/dari data store tidak tercatat (karena tidak bernama di DFD), padahal Modul 5 §1.11 mensyaratkan seluruh flow dan store tercatat.", "Tinggi",
     "Setelah VS-01 diperbaiki, tambahkan seluruh aliran store ke inventaris dan buat entri kompositnya."),
    ("VS-11", "M05", "Entri 8, 14; penguraian 'rincianDenda', 'bukti Pengembalian'", "Konsistensi nama",
     "Nama entri berbeda dengan inventaris/DFD: laporanBulanan vs LaporanTransaksiBulanan; permintaanValidasi vs dataPermintaanValidasi; "
     "judul penguraian 'rincianDenda' untuk infoDenda; 'bukti Pengembalian' memakai spasi.", "Sedang",
     "Gunakan satu nama baku yang sama persis dengan DFD; nama lain dicatat sebagai alias."),
    ("VS-12", "M05", "Struktur infoDenda, buktiPengembalian, dataPermintaanValidasi, Data Transaksi", "Notasi kunci",
     "@ dipakai ganda pada aliran (@IDTransaksi + @IDAnggota), TanggalPengajuan ditandai sebagai kunci, sedangkan struktur store Data Transaksi "
     "justru tidak memiliki @.", "Sedang",
     "Gunakan @ hanya untuk kunci record data store; di aliran cukup tulis nama elemen."),
    ("VS-13", "M05", "pengajuanPeminjaman, buktiPeminjaman, Data Transaksi", "Konsistensi struktur",
     "Jumlah buku per transaksi tidak konsisten: aturan '1 anggota hanya bisa meminjam 1 buku', buktiPeminjaman {KodeBuku} dengan JudulBuku tunggal, "
     "Data Transaksi 1{KodeBuku}5.", "Tinggi",
     "Tetapkan kardinalitas satu kali (mis. 1{KodeBuku}5) dan terapkan di semua struktur; bila banyak buku, jadikan {KodeBuku + JudulBuku}."),
    ("VS-14", "M05", "Penguraian hari, bulan, tahun, Pengarang, huruf", "Notasi",
     "hari = {1 | 2 | ... | 31} memakai {} (pengulangan) padahal maksudnya pilihan []. Pengarang = 1{huruf}3 berarti 1-3 huruf, bertentangan "
     "dengan tabel elemen 1{huruf}50 dan entri Data Koleksi (1{Pengarang}3). huruf = [A..Z] vs [A..z].", "Sedang",
     "Ganti {} menjadi [] untuk pilihan. Bedakan jumlah pengarang (1{Pengarang}3) dari panjang nama pengarang (1{karakter}50). Satukan definisi huruf."),
    ("VS-15", "M05", "StatusAnggota, StatusKeanggotaan, IDPeminjaman, StatusPeminjaman, KategoriAnggota", "Sinonim/homonim",
     "Satu konsep punya beberapa nama/domain: StatusAnggota [Aktif | NonAktif] vs StatusKeanggotaan [Aktif | Non-Aktif]; IDTransaksi vs IDPeminjaman; "
     "KategoriAnggota [mahasiswa | dosen] vs [Mahasiswa | Dosen]; StatusPeminjaman [Tersedia | Tidak Tersedia] memakai domain status buku. "
     "Hubungan NoIdentitas (9-18 digit) dengan IDAnggota (10 digit) tidak dijelaskan.", "Tinggi",
     "Pilih satu nama baku per konsep, catat sisanya di daftar alias, dan jelaskan cara IDAnggota diturunkan dari NoIdentitas."),
    ("VS-16", "M05", "NamaDepan, NamaBelakang, TotalDenda, tahun, Hari", "Domain bertentangan",
     "NamaDepan 1{huruf}20 vs 1{huruf}30; NamaBelakang wajib vs opsional; TotalDenda 1{angka}6 vs 1{angka}8; tahun 2023-2500 vs 2025-2500 vs 4{angka}4; "
     "Hari [1..31] vs 2{angka}2.", "Sedang",
     "Tetapkan satu definisi per elemen di tabel elemen, lalu rujuk definisi itu dari semua struktur."),
    ("VS-17", "M05", "Tabel elemen: NoTelepon, EkstensiDomain, TanggalJatuhTempo", "Contoh vs aturan",
     "NoTelepon: contoh valid 81234567890 padahal aturan 'diawali 08'. EkstensiDomain: contoh 'ac.id' mengandung titik padahal domain hanya huruf. "
     "TanggalJatuhTempo: aturan 7 hari kerja, tetapi contoh 19-10 -> 26-10 adalah 7 hari kalender (7 hari kerja = 28-10-2026).", "Sedang",
     "Samakan contoh dengan aturan, atau ubah aturannya."),
    ("VS-18", "M05", "Atribut Volume semua entri; daftar TBD & alias", "Kelengkapan",
     "Volume diisi 'TBD' tanpa pemilik dan batas waktu. Dokumen tidak memuat daftar TBD dan daftar alias (isi minimum dokumen utama Modul 5 Bagian IV).", "Sedang",
     "Isi volume sebagai estimasi dengan asumsi, atau buat daftar TBD dengan pemilik dan tenggat; tambahkan daftar alias."),
    ("VS-19", "M05", "Entri 15-17 (data store)", "Kelengkapan",
     "Entri store belum menyebut kunci, retensi, dan akses secara eksplisit (Modul 5 §1.3). Alias 'Tabel Anggota' adalah istilah fisik.", "Rendah",
     "Tambahkan atribut kunci, retensi, sumber, akses; gunakan istilah bisnis untuk alias."),
    ("VS-20", "M05", "Tabel elemen", "Kelengkapan",
     "KodePos, Hari, Bulan, Penerbit, Pengarang tanpa nullability dan sumber nilai. StatusKeanggotaan, IDPeminjaman, StatusPeminjaman, "
     "TarifDendaPerHari, TotalDendaBulanan dipakai di struktur/aturan tetapi tidak ada di tabel elemen.", "Sedang",
     "Lengkapi atribut elemen dan tambahkan elemen yang belum didefinisikan."),
    ("VS-21", "M05", "Entri 12, 13, 14, 17", "Notasi",
     "Struktur ditulis tanpa ruas kiri (mis. '[Terlambat | Tepat waktu]', '@KodeBuku + [...]', 'PeriodeLaporan + ...').", "Rendah",
     "Tulis lengkap 'namaData = struktur'."),
    ("VS-22", "M05", "Entri buktiPengembalian, aturan bisnis", "Semantik",
     "Bukti hanya diterbitkan bila denda lunas, tetapi DFD tidak punya aliran pembayaran denda sehingga sistem tidak tahu kapan denda lunas.", "Sedang",
     "Tambahkan aliran pembayaran/konfirmasi denda di DFD, atau ubah aturannya."),
    ("VS-23", "M05", "Bagian Validasi Silang", "Kelengkapan",
     "Bagian Validasi Silang kosong; tidak ada revision history.", "Sedang",
     "Lakukan Langkah Kerja 5 Modul 5 untuk minimal satu proses dan catat hasilnya."),
]

BAIK = [
    "Penguraian struktur sampai elemen dasar (angka, huruf) sudah dilakukan untuk hampir semua aliran, termasuk Email dan Alamat.",
    "Setiap elemen pada tabel elemen sudah memiliki contoh valid dan contoh tidak valid.",
    "Inventaris aliran yang bernama sudah memuat sumber dan tujuan dengan nomor proses yang jelas.",
    "Aturan bisnis denda dirumuskan sebagai perhitungan (TerlambatHari x tarif), sehingga mudah diturunkan ke process spec.",
]

# --------------------------------------------------------------- Rekomendasi penulisan
REKOMENDASI = [
    # aspek, acuan, sekarang (dokumen K4), disarankan
    ("1. Nama entri sama persis dengan DFD", "Modul 5 Langkah 1.2, §1.8",
     "Inventaris: laporanBulanan; entri komposit: LaporanTransaksiBulanan. Entri permintaanValidasi, struktur dataPermintaanValidasi.",
     "Nama = laporanBulanan di DFD, inventaris, entri, dan penguraian. Nama lain masuk kolom Alias."),
    ("2. Semua aliran data store diberi nama lalu diinventarisasi", "Modul 1-3 tabel komponen DFD; Modul 5 §1.11",
     "D1 -> 2.1 tanpa nama, tidak ada entri.",
     "Beri nama di DFD, mis. D1 -> 2.1: dataAnggotaAktif = IDAnggota + StatusAnggota. Tambahkan ke inventaris (Jenis: Flow) dan buat entri kompositnya."),
    ("3. Simbol @ hanya untuk kunci record data store", "Modul 5 §1.4, Langkah 4.4",
     "infoDenda = @IDTransaksi + @IDAnggota + NamaAnggota + TerlambatHari + TotalDenda",
     "infoDenda = IDTransaksi + IDAnggota + NamaAnggota + TerlambatHari + TotalDenda\nData Transaksi = {RecordTransaksi}; RecordTransaksi = @IDTransaksi + ..."),
    ("4. [] untuk pilihan, {} untuk pengulangan", "Modul 5 §1.4 Gambar 31",
     "hari = { 1 | 2 | ... | 31 }",
     "Hari = 2{angka}2 (domain 01-31 ditulis di tabel elemen), atau Hari = [01 | 02 | ... | 31]"),
    ("5. Satu elemen, satu definisi", "Modul 5 §1.8, §1.11",
     "NamaDepan = 1{huruf}20 di satu struktur, 1{huruf}30 di struktur lain; TotalDenda 1{angka}6 vs 1{angka}8.",
     "Definisikan sekali di tabel elemen, lalu struktur lain cukup merujuk nama elemen. Pengecekan: setiap elemen hanya punya satu baris di tabel elemen."),
    ("6. Kardinalitas konsisten", "Modul 5 Langkah 3.3",
     "pengajuanPeminjaman 1 buku; buktiPeminjaman {KodeBuku} + JudulBuku; Data Transaksi 1{KodeBuku}5.",
     "Tetapkan sekali, mis. 1{DetailBuku}5 dengan DetailBuku = KodeBuku + JudulBuku, lalu pakai di semua struktur."),
    ("7. Contoh valid harus lolos aturan, contoh tidak valid melanggar satu aturan", "Modul 5 §1.7, §1.11",
     "NoTelepon: aturan 'diawali 08', contoh valid 81234567890.",
     "Contoh valid: 081234567890. Contoh tidak valid: 81234567890 (tidak diawali 08). Sebutkan aturan yang dilanggar."),
    ("8. Struktur ditulis lengkap 'nama = struktur'", "Modul 5 §1.4",
     "Struktur: [Terlambat | Tepat waktu]",
     "statusKeterlambatan = [Terlambat | TepatWaktu]"),
    ("9. Volume sebagai estimasi dengan asumsi; TBD punya pemilik & tenggat", "Modul 5 Langkah 2.4, §1.9",
     "Volume: TBD (semua entri).",
     "Volume: ~100 per hari kerja (asumsi: 10% anggota aktif meminjam per minggu). Bila benar-benar belum tahu: TBD-01, pemilik: <nama>, tenggat: sebelum M06."),
    ("10. Entri data store memuat atribut store", "Modul 5 §1.3",
     "Data Anggota: Nama, Alias 'Tabel Anggota', Arus, Penjelasan, Periode, Volume, Struktur.",
     "Tambah: Record, Kunci (@IDAnggota), Retensi, Sumber (proses penulis), Akses (proses pembaca). Alias tanpa istilah fisik ('tabel')."),
    ("11. Elemen yang dipakai wajib ada di tabel elemen", "Modul 5 §1.10, Langkah 5.2",
     "StatusKeanggotaan, IDPeminjaman, StatusPeminjaman, TarifDendaPerHari tidak ada di tabel elemen.",
     "Setelah selesai menulis struktur, daftar semua nama elemen yang muncul lalu cocokkan dengan tabel elemen. Yang tidak ada: tambahkan, atau ganti dengan nama baku."),
    ("12. Bagian Validasi Silang diisi dengan penelusuran satu proses", "Modul 5 Langkah 5",
     "Bagian kosong.",
     "Tabel: Langkah 1-6 Modul 5 (input-output -> entri, elemen terdefinisi, sinonim, kunci store, perbaikan DFD, TBD), seperti bagian 4 dokumen ini."),
]


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.8))
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)

    def tabel(header, rows, size=9, widths=None):
        t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = h
            for r in c.paragraphs[0].runs:
                r.bold = True; r.font.size = Pt(size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        doc.add_paragraph()
        return t

    doc.add_heading("Validasi Silang M04 & M05 Kelompok 4 oleh Kelompok 6", 0)
    tabel(["Komponen", "Keterangan"], IDENTITAS, 10, [5, 20])

    doc.add_heading("1. Peer Review DFD (Modul 4 Langkah Kerja 5)", 1)
    doc.add_paragraph("Interpretasi ditulis hanya berdasarkan diagram (tanpa penjelasan lisan). "
                      "Checklist: Ya = memenuhi, Tidak = ada temuan (lihat ID temuan).")
    tabel(["Diagram", "Interpretasi Tujuan", "Sintaks", "Semantik", "Balancing", "Keterbacaan", "Catatan"],
          M04_REVIEW, 9, [3.5, 7, 1.6, 1.6, 1.8, 2, 8])

    doc.add_heading("2. Matriks Balancing Hasil Pemeriksaan", 1)
    tabel(["Aliran pada Induk", "Arah", "Aliran Ekuivalen pada Anak", "Status", "Catatan"], BALANCING, 9, [6, 2, 6, 3, 8])

    doc.add_heading("3. Checklist Aturan Kualitas Kamus Data (Modul 5 §1.11)", 1)
    tabel(["Aturan", "Terpenuhi", "Bukti / Catatan"], M05_CHECK, 9, [8, 2.5, 14])

    doc.add_heading("4. Validasi Silang Proses 3 Kelola Pengembalian (Modul 5 Langkah Kerja 5)", 1)
    tabel(["Langkah", "Hasil"], TRACE, 9, [7, 18])

    doc.add_heading("5. Daftar Temuan (format Lampiran C)", 1)
    doc.add_paragraph("Severitas: Tinggi = memengaruhi balancing atau keterlacakan ke M06; Sedang = inkonsistensi definisi; "
                      "Rendah = kerapian. Status awal semua temuan: Open (menunggu respons Kelompok 4: Diterima / Ditolak / Perlu klarifikasi).")
    rows = [t[:6] + (t[6], "Open") for t in TEMUAN]
    tabel(["ID", "Artefak", "Lokasi", "Jenis", "Deskripsi Berbasis Bukti", "Severitas", "Rekomendasi", "Status"],
          rows, 8, [1.3, 1.6, 3.3, 2.2, 7.5, 1.6, 6.2, 1.3])

    ring = {}
    for t in TEMUAN:
        ring[t[5]] = ring.get(t[5], 0) + 1
    doc.add_paragraph("Ringkasan: %d temuan (Tinggi %d, Sedang %d, Rendah %d)." %
                      (len(TEMUAN), ring.get("Tinggi", 0), ring.get("Sedang", 0), ring.get("Rendah", 0)))

    doc.add_heading("6. Hal yang Sudah Baik", 1)
    for b in BAIK:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("7. Rekomendasi Penulisan Kamus Data", 1)
    doc.add_paragraph("Bagian ini merangkum cara penulisan yang disarankan agar temuan pada bagian 5 tidak berulang. "
                      "Setiap butir merujuk ketentuan modul dan memakai contoh dari dokumen Kelompok 4.")
    tabel(["Aspek", "Acuan", "Tertulis Sekarang", "Penulisan yang Disarankan"], REKOMENDASI, 8.5, [5, 3.5, 7.5, 9.5])
    doc.add_paragraph("Urutan isi dokumen utama sesuai Modul 5 Bagian IV: Inventaris -> Entri Komposit (aliran, lalu data store) -> "
                      "Penguraian Struktur -> Elemen -> Daftar Alias -> Daftar TBD -> Validasi Silang, ditambah revision history.")

    doc.save(OUT)
    print("Ditulis:", OUT, "|", ring)


if __name__ == "__main__":
    main()
