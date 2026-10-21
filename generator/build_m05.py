# -*- coding: utf-8 -*-
"""Bangun Kamus Data M05 (xlsx + docx draft) dari m05_data.py.

Pakai:  python manajemen/generator/build_m05.py
Hasil:  progress/PRPL_M05_2B_RajinBelajar_DataDictionary.xlsx
        progress/PRPL_M05_2B_RajinBelajar_DataDictionary_draft.docx
"""
import os, re, sys
sys.path.insert(0, os.path.dirname(__file__))
from m05_data import *

# True  = penguraian sampai elemen dasar (angka, huruf, karakter), seperti Kelompok 4.
# False = berhenti di elemen elementer; format & panjang cukup di tabel elemen
#         (minimum yang diwajibkan Modul 5 §1.4 & Langkah 3.5).
SAMPAI_DASAR = True

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT_XLSX = os.path.join(ROOT, "progress", "PRPL_M05_2B_RajinBelajar_DataDictionary.xlsx")
OUT_DOCX = os.path.join(ROOT, "progress", "PRPL_M05_2B_RajinBelajar_DataDictionary_draft.docx")


# ---------------------------------------------------------------- penguraian
def refs(nama):
    """Nama definisi lain yang dirujuk oleh struktur `nama` (enum tidak diurai)."""
    struktur, kind = DEF[nama]
    if kind == "enum":
        return []
    s = re.sub(r'"[^"]*"', " ", struktur)
    out = []
    for tok in re.findall(r"[A-Za-z][A-Za-z0-9]*", s):
        if tok in DEF and tok != nama and tok not in out:
            out.append(tok)
    return out


def uraikan(nama, penuh=False):
    """Urutan penguraian BFS dari `nama` sampai elemen dasar."""
    urut, antri = [], [nama]
    while antri:
        n = antri.pop(0)
        if n in urut:
            continue
        urut.append(n)
        if not (SAMPAI_DASAR or penuh) and DEF[n][1] in ("elemen", "enum") and n != nama:
            continue
        antri.extend(r for r in refs(n) if r not in urut)
    return urut


def label(nama):
    kind = DEF[nama][1]
    return {"komposit": "Komposit", "record": "Komposit", "elemen": "Elementer",
            "enum": "Elementer", "dasar": "Dasar"}[kind]


# ---------------------------------------------------------------- validasi
def validasi():
    masalah = []
    nama_inv = [r[0] for r in INVENTARIS]
    for n in nama_inv:
        if n not in DEF:
            masalah.append(f"Inventaris '{n}' tidak punya struktur")
        if n in URUT_FLOW and n not in KOMPOSIT:
            masalah.append(f"Flow '{n}' tidak punya entri komposit")
    for n in URUT_STORE:
        if n not in STORE:
            masalah.append(f"Store '{n}' tidak punya entri store")
    # token yang tampak seperti nama data tapi tidak terdefinisi
    for n, (s, kind) in DEF.items():
        if kind in ("enum", "dasar"):
            continue
        for tok in re.findall(r"[A-Za-z][A-Za-z0-9]*", re.sub(r'"[^"]*"', " ", s)):
            if tok not in DEF and tok not in ("D1", "D2", "D3", "Data", "Anggota", "Koleksi", "Transaksi", "Fisik"):
                masalah.append(f"'{n}' merujuk '{tok}' yang belum didefinisikan")
    # setiap elemen yang dipakai harus ada di tabel elemen
    dipakai = set()
    for n in nama_inv:
        dipakai.update(x for x in uraikan(n, True) if DEF[x][1] in ("elemen", "enum"))
    for e in sorted(dipakai):
        if e not in ELEMEN:
            masalah.append(f"Elemen '{e}' dipakai tetapi tidak ada di tabel elemen")
    # tidak sirkular (A -> ... -> A)
    for n in DEF:
        for r in refs(n):
            if n in uraikan(r, True):
                masalah.append(f"Struktur sirkular: {n} <-> {r}")
    return masalah, sorted(dipakai)


# ---------------------------------------------------------------- xlsx
def tulis_xlsx(elemen_dipakai):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = Workbook()
    thin = Side(style="thin", color="999999")
    head = PatternFill("solid", fgColor="DDEBF7")

    def sheet(ws, judul, header, rows, widths):
        ws.append([judul]); ws["A1"].font = Font(bold=True, size=12)
        ws.append([f"Kelompok 6 (RajinBelajar) - 2B | Kamus Data SILAPUS v{VERSI} | {TANGGAL}"])
        ws.append([])
        ws.append(header)
        for c in ws[4]:
            c.font = Font(bold=True); c.fill = head
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows(min_row=4):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A5"

    inv_rows = []
    for i, (n, j, t, lok, src, dst, cd, lama) in enumerate(INVENTARIS, 1):
        status = "Baku" if (lama in (None, n)) else f"Baku (dulu: {lama})"
        inv_rows.append((i, n, j, "Tipe " + str(t), lok, src, dst, cd or "-", status))
    sheet(wb.active, "INVENTARIS DATA (Modul 5 Langkah 1)",
          ["No", "Nama Data", "Jenis", "Tipe", "Lokasi DFD", "Sumber", "Tujuan", "Nama di Context Diagram", "Status Entri"],
          inv_rows, [5, 28, 8, 8, 16, 30, 30, 30, 32])
    wb.active.title = "1_Inventaris"

    kom_rows = []
    for n in URUT_FLOW:
        a = KOMPOSIT[n]
        kom_rows.append((n, a[0], a[1], a[2], a[3], a[4], a[5], f"{n} = {DEF[n][0]}", a[6]))
    sheet(wb.create_sheet("2_Entri_Komposit"), "ENTRI KOMPOSIT ALIRAN DATA (Modul 5 §1.6)",
          ["Nama", "Alias", "Bentuk Data", "Arus Data", "Penjelasan", "Periode", "Volume", "Struktur", "Aturan Bisnis"],
          kom_rows, [24, 26, 16, 30, 40, 26, 26, 50, 50])

    st_rows = []
    for n in URUT_STORE:
        s = STORE[n]; v = STORE_VOL[n]
        rec = s[1]
        st_rows.append((n, s[0], f"{n} = {DEF[n][0]}\n{rec} = {DEF[rec][0]}", s[2], s[3], s[4], s[5], v[0], v[1]))
    sheet(wb.create_sheet("3_Entri_Store"), "ENTRI DATA STORE (Modul 5 §1.3)",
          ["Nama", "Deskripsi", "Record", "Kunci", "Retensi", "Sumber (penulis)", "Akses (pembaca)", "Periode", "Volume"],
          st_rows, [20, 34, 60, 16, 16, 34, 34, 24, 24])

    ur_rows = []
    for n in URUT_FLOW + URUT_STORE:
        for x in uraikan(n):
            ur_rows.append((n, x, DEF[x][0], label(x)))
    sheet(wb.create_sheet("4_Penguraian"), "PENGURAIAN STRUKTUR SAMPAI ELEMEN ATOMIK (Modul 5 Langkah 3)",
          ["Data Induk", "Nama", "Struktur", "Tingkat"], ur_rows, [26, 24, 70, 12])

    el_rows = []
    for e in elemen_dipakai:
        a = ELEMEN[e]
        el_rows.append((e, a[0], a[1], a[2], a[3], a[4], a[5], a[6], a[7]))
    sheet(wb.create_sheet("5_Elemen"), "ENTRI ELEMEN DATA (Modul 5 §1.7 & Langkah 4)",
          ["Nama", "Definisi", "Tipe Konseptual", "Format / Panjang", "Domain / Aturan", "Nullability", "Sumber Nilai", "Contoh Valid", "Contoh Tidak Valid"],
          el_rows, [22, 40, 16, 26, 40, 22, 26, 22, 26])

    sheet(wb.create_sheet("6_Alias"), "DAFTAR ALIAS (Modul 5 §1.8)",
          ["Nama Baku", "Alias / Istilah Lain", "Catatan"], ALIAS, [26, 60, 50])
    sheet(wb.create_sheet("7_TBD"), "DAFTAR TBD (Modul 5 §1.9)",
          ["ID", "Topik", "Hal yang Belum Pasti", "Pemilik", "Batas Waktu"], TBD, [8, 26, 70, 14, 26])
    sheet(wb.create_sheet("8_Parameter"), "PARAMETER ATURAN BISNIS",
          ["Parameter", "Nilai", "Sumber"], PARAMETER, [24, 34, 50])
    wb.save(OUT_XLSX)


# ---------------------------------------------------------------- docx
def tulis_docx(elemen_dipakai):
    from docx import Document
    from docx.shared import Pt, Cm
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)

    def tabel(header, rows, size=9):
        t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
        for i, h in enumerate(header):
            t.rows[0].cells[i].text = h
            for r in t.rows[0].cells[i].paragraphs[0].runs:
                r.bold = True; r.font.size = Pt(size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
        return t

    doc.add_heading("DRAFT v%s - Kamus Data SILAPUS" % VERSI, 0)
    doc.add_paragraph("Draf hasil evaluasi %s. Isi dibangkitkan dari manajemen/generator/m05_data.py; "
                      "salin ke dokumen utama setelah disepakati kelompok." % TANGGAL)

    doc.add_heading("Membuat Inventaris Data", 1)
    rows = []
    for i, (n, j, t, lok, src, dst, cd, lama) in enumerate(INVENTARIS, 1):
        rows.append((i, n, j, lok, src, dst, "Baku" if lama in (None, n) else f"Baku (dulu: {lama})"))
    tabel(["No", "Nama Data", "Jenis", "Lokasi DFD", "Sumber", "Tujuan", "Status Entri"], rows, 8)

    doc.add_heading("Menulis Entri Komposit", 1)
    for i, n in enumerate(URUT_FLOW, 1):
        a = KOMPOSIT[n]
        doc.add_paragraph(f"{i}. {n}", style="List Number" if False else None).runs[0].bold = True
        tabel(["Atribut", "Isi"], [("Nama", n), ("Alias", a[0]), ("Bentuk Data", a[1]), ("Arus Data", a[2]),
                                   ("Penjelasan", a[3]), ("Periode", a[4]), ("Volume", a[5]),
                                   ("Struktur", f"{n} = {DEF[n][0]}"), ("Aturan Bisnis", a[6])])
    doc.add_heading("Entri Data Store", 2)
    for n in URUT_STORE:
        s = STORE[n]; v = STORE_VOL[n]; rec = s[1]
        doc.add_paragraph(n).runs[0].bold = True
        tabel(["Atribut", "Isi"], [("Nama", n), ("Deskripsi", s[0]),
                                   ("Record", f"{n} = {DEF[n][0]}\n{rec} = {DEF[rec][0]}"), ("Kunci", s[2]),
                                   ("Retensi", s[3]), ("Sumber", s[4]), ("Akses", s[5]), ("Periode", v[0]), ("Volume", v[1])])

    doc.add_heading("Menguraikan Struktur Data", 1)
    doc.add_paragraph("Notasi: = didefinisikan sebagai, + dan, () opsional, {} berulang (n{x}m = minimal n maksimal m), "
                      "[] pilih satu, | pemisah alternatif, @ kunci. Setiap struktur diuraikan sampai elemen dasar.")
    for n in URUT_FLOW + URUT_STORE:
        doc.add_paragraph(n).runs[0].bold = True
        for x in uraikan(n):
            p = doc.add_paragraph(f"{x} = {DEF[x][0]}   ({label(x)})")
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(9.5); r.font.name = "Consolas"

    doc.add_heading("Mendefinisikan Elemen", 1)
    rows = [(e,) + ELEMEN[e] for e in elemen_dipakai]
    tabel(["Nama", "Definisi", "Tipe", "Format/Panjang", "Domain/Aturan", "Null", "Sumber Nilai", "Contoh Valid", "Contoh Tidak Valid"], rows, 7.5)

    doc.add_heading("Daftar Alias", 1)
    tabel(["Nama Baku", "Alias", "Catatan"], ALIAS)
    doc.add_heading("Daftar TBD", 1)
    tabel(["ID", "Topik", "Hal yang Belum Pasti", "Pemilik", "Batas Waktu"], TBD)
    doc.add_heading("Parameter Aturan Bisnis", 1)
    tabel(["Parameter", "Nilai", "Sumber"], PARAMETER)
    doc.add_heading("Validasi Silang", 1)
    doc.add_paragraph("Lihat manajemen/EVALUASI_DATA_DICTIONARY.md bagian 'Validasi Silang Internal' (proses 3.1).")
    doc.save(OUT_DOCX)


if __name__ == "__main__":
    masalah, dipakai = validasi()
    if masalah:
        print("MASALAH:"); [print(" -", m) for m in masalah]
        sys.exit(1)
    print(f"OK: {len(INVENTARIS)} entri inventaris, {len(KOMPOSIT)} komposit, {len(STORE)} store, {len(dipakai)} elemen")
    tulis_xlsx(dipakai); tulis_docx(dipakai)
    print("Ditulis:", OUT_XLSX); print("Ditulis:", OUT_DOCX)
